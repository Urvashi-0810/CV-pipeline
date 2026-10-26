"""
STEP 2c - Store candidates in Excel (pandas + openpyxl) and read them back.

One row per CV file. The "File Hash" column is the row's unique key, so
re-extracting a CV (--force) replaces its row instead of adding a duplicate.

The shortlisting step (next part) will call `load_candidates()` to get a DataFrame.
"""
from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.logger import get_logger
from app.core.models import Candidate

log = get_logger("excel")

SHEET = "Candidates"
KEY = "File Hash"

# Column order in the Excel file (and width in characters)
COLUMNS: dict[str, int] = {
    "Full Name": 24, "Email": 30, "Phone": 18, "Passport Number": 20, "Skills": 50,
    "Experience (Years)": 12, "Education": 50, "Current/Last Job Title": 28,
    "Current/Last Company": 24, "Age": 6, "Date of Birth": 13,
    "Certifications": 35, "Languages": 22, "Location": 22, "Summary": 60,
    "CV File": 40, "Received From": 30, "Mail Subject": 30, "Mail Date": 18,
    "Extraction Method": 14, "LLM": 10, "Extracted At": 18, "Uploaded to S3": 18, KEY: 20,
}


def _fmt_education(c: Candidate) -> str:
    parts = []
    for e in c.education:
        s = e.degree or ""
        if e.institution:
            s += f" - {e.institution}"
        if e.year:
            s += f" ({e.year})"
        if e.score:
            s += f" [{e.score}]"
        if s.strip():
            parts.append(s.strip(" -"))
    return "; ".join(parts)


def candidate_to_row(c: Candidate, *, file_name: str, file_hash: str, mail_meta: dict,
                     method: str, provider: str) -> dict:
    return {
        "Full Name": c.full_name,
        "Email": c.email,
        "Phone": c.phone,
        "Passport Number": c.passport_number,
        "Skills": ", ".join(c.skills),
        "Experience (Years)": c.total_experience_years,
        "Education": _fmt_education(c),
        "Current/Last Job Title": c.current_job_title,
        "Current/Last Company": c.current_company,
        "Age": c.age,
        "Date of Birth": c.date_of_birth,
        "Certifications": ", ".join(c.certifications),
        "Languages": ", ".join(c.languages),
        "Location": c.location,
        "Summary": c.summary,
        "CV File": file_name,
        "Received From": mail_meta.get("from"),
        "Mail Subject": mail_meta.get("subject"),
        "Mail Date": (mail_meta.get("mail_date") or "")[:16].replace("T", " ") or None,
        "Extraction Method": method,
        "LLM": provider,
        "Extracted At": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Uploaded to S3": "Pending",
        KEY: file_hash,
    }


def load_candidates(path: Path | None = None) -> pd.DataFrame:
    path = path or settings.output_excel
    if not path.exists():
        return pd.DataFrame(columns=list(COLUMNS))
    df = pd.read_excel(path, sheet_name=SHEET, dtype={KEY: str, "Phone": str})
    return df


def upsert_rows(rows: list[dict], path: Path | None = None) -> int:
    """Insert new rows / replace rows with the same File Hash. Returns total row count."""
    if not rows:
        return len(load_candidates(path))
    path = path or settings.output_excel
    df = pd.concat([load_candidates(path), pd.DataFrame(rows)], ignore_index=True)
    df = df.drop_duplicates(subset=[KEY], keep="last")
    df = df.reindex(columns=list(COLUMNS))
    _write(df, path)
    return len(df)


def _write(df: pd.DataFrame, path: Path) -> None:
    write_sheets(path, {SHEET: df}, widths=COLUMNS)


def write_sheets(path: Path, sheets: dict[str, pd.DataFrame], widths: dict[str, int] | None = None) -> None:
    """
    Write one or more DataFrames as formatted sheets (blue header, frozen first row/column,
    filters). Used by every step (candidates, shortlist, job-fit) so they all look the same.
    Writes to a temp file first, then swaps it in; retries if the file is open in Excel
    (Windows locks open files).
    """
    widths = widths or {}
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for i, col in enumerate(df.columns, start=1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(i)].width = widths.get(col, COLUMNS.get(col, 18))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=False, vertical="top")
            ws.freeze_panes = "B2"
            if len(df.columns):
                ws.auto_filter.ref = ws.dimensions

    for attempt in range(5):
        try:
            tmp.replace(path)
            return
        except PermissionError:
            log.warning("%s is open in another program - close it! retrying (%d/5)...", path.name, attempt + 1)
            time.sleep(3)
    raise PermissionError(f"Could not write {path}. Close it in Excel; new data is in {tmp}")
